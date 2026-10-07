"""Build a multi-sheet Excel operational report aligned to the Power BI pages."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.analytics import (
    compute_kpis,
    department_performance,
    insurance_breakdown,
    load_patients,
    monthly_admissions,
    top_doctors,
)
from src.config import EXCEL_REPORT, ensure_directories

HEADER_FILL = PatternFill("solid", fgColor="0F4C81")
HEADER_FONT = Font(color="FFFFFF", bold=True)
KPI_FILL = PatternFill("solid", fgColor="E8F1FA")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN


def _write_table(ws, df, start_row: int, table_name: str) -> int:
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start_row):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.border = THIN
            if r_idx == start_row:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
    end_row = start_row + len(df)
    end_col = chr(64 + max(len(df.columns), 1))
    table = Table(displayName=table_name, ref=f"A{start_row}:{end_col}{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    for column in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max(max_length + 2, 14), 36)
    return end_row


def _kpi_block(ws, kpis: dict[str, float]) -> None:
    labels = [
        ("Total Patients", f"{kpis['total_patients']:,}"),
        ("Total Revenue", f"${kpis['total_revenue']:,.2f}"),
        ("Avg Billing", f"${kpis['avg_billing']:,.2f}"),
        ("Avg Length of Stay", f"{kpis['avg_length_of_stay']:.1f} days"),
    ]
    for idx, (label, value) in enumerate(labels, start=1):
        ws.cell(1, idx, label).fill = HEADER_FILL
        ws.cell(1, idx).font = HEADER_FONT
        ws.cell(2, idx, value).fill = KPI_FILL
        ws.cell(2, idx).font = Font(bold=True, size=14)
        ws.column_dimensions[chr(64 + idx)].width = 24


def generate_excel_report(output_path: Path | None = None) -> Path:
    ensure_directories()
    df = load_patients()
    kpis = compute_kpis(df)
    dept = department_performance(df)
    monthly = monthly_admissions(df)
    doctors = top_doctors(df, n=15)
    insurance = insurance_breakdown(df)
    output = output_path or EXCEL_REPORT

    wb = Workbook()

    exec_ws = wb.active
    exec_ws.title = "Executive"
    _kpi_block(exec_ws, kpis)
    exec_ws["A4"] = "Department performance"
    exec_ws["A4"].font = Font(bold=True, size=12, color="0F4C81")
    _write_table(exec_ws, dept, start_row=5, table_name="ExecutiveDepartments")
    chart = BarChart()
    chart.title = "Revenue by Department"
    chart.y_axis.title = "Revenue"
    data = Reference(exec_ws, min_col=3, min_row=5, max_row=5 + len(dept))
    cats = Reference(exec_ws, min_col=1, min_row=6, max_row=5 + len(dept))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    exec_ws.add_chart(chart, "G5")

    adm_ws = wb.create_sheet("Admissions")
    adm_ws["A1"] = "Monthly admissions and revenue"
    adm_ws["A1"].font = Font(bold=True, size=12, color="0F4C81")
    adm_view = monthly[["Period", "Admissions", "Revenue"]]
    _write_table(adm_ws, adm_view, start_row=3, table_name="MonthlyAdmissions")
    line = LineChart()
    line.title = "Monthly Admissions"
    line.y_axis.title = "Admissions"
    line.add_data(Reference(adm_ws, min_col=2, min_row=3, max_row=3 + len(adm_view)), titles_from_data=True)
    line.set_categories(Reference(adm_ws, min_col=1, min_row=4, max_row=3 + len(adm_view)))
    adm_ws.add_chart(line, "E3")

    doc_ws = wb.create_sheet("Doctors")
    doc_ws["A1"] = "Top doctors by patient volume"
    doc_ws["A1"].font = Font(bold=True, size=12, color="0F4C81")
    _write_table(doc_ws, doctors, start_row=3, table_name="TopDoctors")
    bar = BarChart()
    bar.type = "bar"
    bar.title = "Patients handled"
    bar.add_data(Reference(doc_ws, min_col=2, min_row=3, max_row=3 + len(doctors)), titles_from_data=True)
    bar.set_categories(Reference(doc_ws, min_col=1, min_row=4, max_row=3 + len(doctors)))
    doc_ws.add_chart(bar, "E3")

    fin_ws = wb.create_sheet("Financial")
    fin_ws["A1"] = "Insurance and department financials"
    fin_ws["A1"].font = Font(bold=True, size=12, color="0F4C81")
    _write_table(fin_ws, insurance, start_row=3, table_name="InsuranceFinancials")
    pie = PieChart()
    pie.title = "Billed amount by insurer"
    pie.add_data(Reference(fin_ws, min_col=3, min_row=3, max_row=3 + len(insurance)), titles_from_data=True)
    pie.set_categories(Reference(fin_ws, min_col=1, min_row=4, max_row=3 + len(insurance)))
    fin_ws.add_chart(pie, "E3")
    fin_ws["A" + str(6 + len(insurance))] = "Department revenue"
    _write_table(
        fin_ws,
        dept[["Department", "TotalRevenue", "AvgBilling"]],
        start_row=7 + len(insurance),
        table_name="DepartmentFinancials",
    )

    patients_ws = wb.create_sheet("Patients")
    export_cols = [
        "PatientID",
        "Age",
        "Gender",
        "Department",
        "Doctor",
        "MedicalCondition",
        "AdmissionType",
        "DateOfAdmission",
        "DischargeDate",
        "LengthOfStay",
        "BillingAmount",
        "InsuranceProvider",
        "Hospital",
        "TestResults",
    ]
    available = [col for col in export_cols if col in df.columns]
    patient_export = df[available].copy()
    for col in ("DateOfAdmission", "DischargeDate"):
        if col in patient_export.columns:
            patient_export[col] = patient_export[col].dt.strftime("%Y-%m-%d")
    _write_table(patients_ws, patient_export, start_row=1, table_name="PatientsExport")

    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(output)
    except PermissionError:
        fallback = output.with_name(f"{output.stem}_latest{output.suffix}")
        wb.save(fallback)
        print(f"Primary workbook is locked; wrote {fallback} instead.")
        return fallback
    return output


if __name__ == "__main__":
    path = generate_excel_report()
    print(f"Wrote {path}")
